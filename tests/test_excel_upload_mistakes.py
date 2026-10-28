"""What happens to a downloaded workbook in Excel before it comes back as an upload.

People retype headings, drag columns about, type a title above the table, rename
the sheet, sort, paste, copy a row, leave a workbook sitting while a teammate
works in GREMLIN, and type "pm reset" for PM_RESET. Each test does one of those
to a real download, saved back by openpyxl the way a spreadsheet program saves
it, and checks the upload reads what was meant or refuses saying why. And the
download itself makes the worst of them hard to do by accident.
"""

import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

import openpyxl

from services import disposition_standard as std
from services.life_data_service import EXCEL_STYLE_UNLOCKED
from services.standard_disposition_service import EXCEL_VERSION_COLUMN
from tests.test_standard_disposition_service import PFS, S09, StandardDispositionFixture, coded

NS = {"main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


class DispositionWorkbookFixture(StandardDispositionFixture):
    def setUp(self) -> None:
        super().setUp()
        self.add_record(501, PFS, "Oven filter clogged", "2024-01-01", "2024-01-01")
        self.add_record(502, PFS, "Belt squeal", "2024-01-02", "2024-01-02")
        self.add_record(503, PFS, "Guard loose", "2024-01-03", "2024-01-03")
        self.std.save_dispositions([{"mapped_record_id": 501, **coded()}], decided_by="tim")
        folder = tempfile.TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        self.path = Path(folder.name) / "dispositions.xlsx"

    def download(self):
        self.std.export_excel(PFS, self.path)
        workbook = openpyxl.load_workbook(self.path)
        return workbook, workbook["Dispositions"]

    @staticmethod
    def column(sheet, heading, header_row=1):
        return next(cell.column for cell in sheet[header_row] if cell.value == heading)

    @staticmethod
    def row(sheet, mapped_record_id):
        return next(cell.row for cell in sheet["A"] if cell.value == mapped_record_id)

    def fill(self, sheet, mapped_record_id, **values):
        row = self.row(sheet, mapped_record_id)
        for heading, value in values.items():
            sheet.cell(row, self.column(sheet, heading)).value = value

    def upload(self, workbook):
        workbook.save(self.path)
        return self.std.import_excel(PFS, self.path, decided_by="jose")

    def saved(self, mapped_record_id):
        return self.std.disposition_record(mapped_record_id)["disposition"]

    def belt_failure(self):
        return coded(component="Belt", mechanism="WEAR", mode="VIB", action="REPLACE")


class TheDownloadTests(DispositionWorkbookFixture):
    def test_its_headings_and_columns_are_locked_and_its_cells_are_not(self):
        self.std.export_excel(PFS, self.path)
        with zipfile.ZipFile(self.path) as workbook:
            sheet = ET.fromstring(workbook.read("xl/worksheets/sheet1.xml"))
            styles = ET.fromstring(workbook.read("xl/styles.xml"))
        protection = sheet.find("main:sheetProtection", NS)
        self.assertEqual(protection.get("sheet"), "1")
        # Rows can be sorted, filtered, added and deleted; columns can't be added
        # or deleted (left at a protected sheet's default), and there is no password.
        for allowed in ("sort", "autoFilter", "insertRows", "deleteRows", "formatColumns"):
            self.assertEqual(protection.get(allowed), "0", allowed)
        for locked in ("insertColumns", "deleteColumns", "password"):
            self.assertIsNone(protection.get(locked), locked)
        xfs = styles.find("main:cellXfs", NS).findall("main:xf", NS)

        def locked(style):
            node = xfs[int(style or 0)].find("main:protection", NS)
            return node is None or node.get("locked") != "0"

        rows = sheet.findall("main:sheetData/main:row", NS)
        self.assertTrue(all(locked(cell.get("s")) for cell in rows[0]))
        self.assertFalse(any(locked(cell.get("s")) for row in rows[1:] for cell in row))
        # A row typed below the table is unlocked too: every column defaults to the unlocked style.
        columns = sheet.findall("main:cols/main:col", NS)
        self.assertTrue(all(column.get("style") == str(EXCEL_STYLE_UNLOCKED) for column in columns))

    def test_the_version_column_is_last_and_hidden(self):
        workbook, sheet = self.download()
        last = sheet.max_column
        self.assertEqual(sheet.cell(1, last).value, EXCEL_VERSION_COLUMN)
        self.assertTrue(sheet.column_dimensions[openpyxl.utils.get_column_letter(last)].hidden)
        self.assertIsNotNone(sheet.cell(self.row(sheet, 501), last).value)
        self.assertIsNone(sheet.cell(self.row(sheet, 502), last).value)


class HeadingsAndColumnsTests(DispositionWorkbookFixture):
    def test_retyped_and_moved_headings_are_read_as_their_columns(self):
        workbook, sheet = self.download()
        self.fill(sheet, 502, **self.belt_failure())
        # The headings the page uses, in other capitals and spacing.
        renames = {"component_code": "Component", "failure_mechanism_code": "Failure Mechanism",
                   "disposition_category": "category", "sub_asset": "Sub-asset", "action_code": "ACTION"}
        for cell in sheet[1]:
            cell.value = renames.get(cell.value, cell.value)
        # And the component column dragged to the front.
        component = self.column(sheet, "Component")
        sheet.insert_cols(1)
        for row in range(1, sheet.max_row + 1):
            sheet.cell(row, 1).value = sheet.cell(row, component + 1).value
            sheet.cell(row, component + 1).value = None
        sheet.cell(1, component + 1).value = "spare"
        outcome = self.upload(workbook)
        self.assertEqual(outcome["saved"], 1)
        self.assertEqual(self.saved(502)["component_code"], "Belt")
        self.assertEqual(outcome["columns_not_read"], ["spare"])

    def test_the_table_is_found_under_a_title_on_a_renamed_sheet_behind_another(self):
        workbook, sheet = self.download()
        self.fill(sheet, 502, **self.belt_failure())
        sheet.insert_rows(1, amount=2)
        sheet.cell(1, 1).value = "PFS backlog, week 43"
        sheet.title = "Week 43"
        workbook.create_sheet("Notes", 0).cell(1, 1).value = "Things to ask the planner"
        self.assertEqual(self.upload(workbook)["saved"], 1)
        self.assertEqual(self.saved(502)["component_code"], "Belt")

    def test_a_missing_column_keeps_what_is_saved_and_is_reported(self):
        workbook, sheet = self.download()
        sheet.delete_cols(self.column(sheet, "disposition_note"))
        self.fill(sheet, 502, **self.belt_failure())
        outcome = self.upload(workbook)
        self.assertEqual(outcome["saved"], 1)
        self.assertEqual(outcome["columns_missing"], ["disposition_note"])

    def test_a_column_twice_refuses_the_workbook_naming_both(self):
        workbook, sheet = self.download()
        extra = sheet.max_column + 1
        sheet.cell(1, extra).value = "Component"
        with self.assertRaisesRegex(ValueError, r"^Columns \w+ and \w+ are both component_code\. Delete the extra one"):
            self.upload(workbook)

    def test_a_workbook_from_before_the_version_column_still_uploads(self):
        workbook, sheet = self.download()
        sheet.delete_cols(sheet.max_column)
        self.fill(sheet, 502, **self.belt_failure())
        self.assertEqual(self.upload(workbook)["saved"], 1)

    def test_the_code_lists_workbook_is_named_as_such(self):
        self.std.export_code_lists_excel(self.path)
        with self.assertRaisesRegex(ValueError, r"^This is the code lists workbook\. Upload it on Developer > Code lists"):
            self.std.import_excel(PFS, self.path)


class TypedValuesTests(DispositionWorkbookFixture):
    def test_codes_read_in_any_capitals_and_by_their_names(self):
        workbook, sheet = self.download()
        codebook = self.std.codebook()
        self.fill(
            sheet, 502,
            disposition_category="failure",
            sub_asset=S09.lower(),
            component_code="belt",
            failure_mechanism_code=codebook.name(std.FAILURE_MECHANISM, "WEAR"),
            failure_mode_code=f"VIB · {codebook.name(std.FAILURE_MODE, 'VIB')}",
            action_code="replace",
            evidence="Explicit",
            renewed_items="filter",
        )
        self.assertEqual(self.upload(workbook)["saved"], 1)
        saved = self.saved(502)
        self.assertEqual(
            (saved["disposition_category"], saved["sub_asset"], saved["component_code"], saved["failure_mechanism_code"],
             saved["failure_mode_code"], saved["action_code"], saved["evidence"], saved["renewed_items"]),
            ("FAILURE", S09, "Belt", "WEAR", "VIB", "REPLACE", "EXPLICIT", ["Filter"]),
        )

    def test_a_category_label_reads_as_its_code_and_other_capitals_change_nothing(self):
        workbook, sheet = self.download()
        # 501 retyped as it is saved, in other capitals: nothing to save.
        self.fill(sheet, 501, disposition_category="Failure", component_code="FILTER")
        self.fill(sheet, 503, disposition_category="Not reviewed")
        outcome = self.upload(workbook)
        self.assertEqual((outcome["saved"], outcome["unchanged"]), (0, 3))


class RowsTests(DispositionWorkbookFixture):
    def test_a_row_without_its_record_id_is_found_by_its_task(self):
        workbook, sheet = self.download()
        self.fill(sheet, 502, **self.belt_failure())
        sheet.cell(self.row(sheet, 502), 1).value = None
        self.assertEqual(self.upload(workbook)["saved"], 1)
        self.assertEqual(self.saved(502)["component_code"], "Belt")

    def test_rows_out_of_line_are_caught(self):
        # Sorting only some columns leaves a row's record id beside another's task.
        workbook, sheet = self.download()
        task = self.column(sheet, "task_id")
        first, second = self.row(sheet, 502), self.row(sheet, 503)
        sheet.cell(first, task).value, sheet.cell(second, task).value = (
            sheet.cell(second, task).value, sheet.cell(first, task).value
        )
        self.fill(sheet, 502, **self.belt_failure())
        with self.assertRaisesRegex(ValueError, r"mapped_record_id 502 is task 502, but the row says task 503"):
            self.upload(workbook)
        self.assertIsNone(self.saved(502)["component_code"])

    def test_a_disposition_with_nothing_to_find_its_work_order_by_is_refused(self):
        workbook, sheet = self.download()
        new_row = sheet.max_row + 1
        sheet.cell(new_row, self.column(sheet, "disposition_category")).value = "FAILURE"
        with self.assertRaisesRegex(ValueError, rf"^Row {new_row}: it has a disposition but no mapped_record_id or task_id"):
            self.upload(workbook)

    def test_a_copied_row_is_refused(self):
        workbook, sheet = self.download()
        source = self.row(sheet, 502)
        copy = sheet.max_row + 1
        for column in range(1, sheet.max_column + 1):
            sheet.cell(copy, column).value = sheet.cell(source, column).value
        with self.assertRaisesRegex(ValueError, rf"^Row {copy}: task 502 is in the workbook twice, also on row {source}\."):
            self.upload(workbook)

    def test_every_refused_row_is_named_at_once_by_the_row_excel_shows(self):
        workbook, sheet = self.download()
        sheet.insert_rows(self.row(sheet, 503))  # an empty row left in the table
        self.fill(sheet, 502, disposition_category="FAILURE")  # FAILURE needs its codes
        orphan = sheet.max_row + 1
        sheet.cell(orphan, self.column(sheet, "component_code")).value = "Belt"
        row_502 = self.row(sheet, 502)
        with self.assertRaises(ValueError) as refused:
            self.upload(workbook)
        lines = str(refused.exception).split("\n")
        self.assertEqual(lines[0], "2 rows were refused, so nothing was saved:")
        self.assertTrue(lines[1].startswith(f"Row {row_502} (task 502): FAILURE needs"), lines[1])
        self.assertTrue(lines[2].startswith(f"Row {orphan}: it has a disposition"), lines[2])


class StaleWorkbookTests(DispositionWorkbookFixture):
    def test_a_row_nobody_touched_keeps_a_teammates_newer_disposition(self):
        workbook, sheet = self.download()
        # After the download, a teammate recodes 501 and dispositions 503 on the page.
        self.std.save_dispositions([
            {"mapped_record_id": 501, **coded(mechanism="PLUG")},
            {"mapped_record_id": 503, **coded("NON_FAILURE", mode="NFF", action="INSPECT_ONLY")},
        ], decided_by="tim")
        self.fill(sheet, 502, **self.belt_failure())
        outcome = self.upload(workbook)
        self.assertEqual(outcome["saved"], 1)
        self.assertEqual(sorted(outcome["kept_newer"]), ["501", "503"])
        self.assertEqual(self.saved(501)["failure_mechanism_code"], "PLUG")
        self.assertEqual(self.saved(503)["disposition_category"], "NON_FAILURE")
        self.assertEqual(self.saved(502)["component_code"], "Belt")

    def test_a_row_changed_in_both_places_is_refused_not_overwritten(self):
        workbook, sheet = self.download()
        self.std.save_dispositions([{"mapped_record_id": 501, **coded(mechanism="PLUG")}], decided_by="tim")
        self.fill(sheet, 501, failure_mechanism_code="WEAR")
        self.fill(sheet, 502, **self.belt_failure())
        with self.assertRaisesRegex(
            ValueError, r"^Row \d+ \(task 501\): tim changed this work order's disposition in GREMLIN on .*after this "
            r"workbook was downloaded, and the row changes it too\. Delete the row to keep what GREMLIN has"
        ):
            self.upload(workbook)
        self.assertEqual(self.saved(501)["failure_mechanism_code"], "PLUG")
        self.assertIsNone(self.saved(502)["component_code"])

    def test_a_row_made_to_agree_with_the_newer_one_is_simply_unchanged(self):
        workbook, sheet = self.download()
        self.std.save_dispositions([{"mapped_record_id": 501, **coded(mechanism="PLUG")}], decided_by="tim")
        self.fill(sheet, 501, failure_mechanism_code="PLUG")
        outcome = self.upload(workbook)
        self.assertEqual((outcome["saved"], outcome["unchanged"], outcome["kept_newer"]), (0, 3, []))


class CodeListWorkbookMistakeTests(StandardDispositionFixture):
    def setUp(self) -> None:
        super().setUp()
        folder = tempfile.TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        self.path = Path(folder.name) / "codes.xlsx"
        self.std.export_code_lists_excel(self.path)
        self.workbook = openpyxl.load_workbook(self.path)
        self.sheet = self.workbook["Code lists"]

    def upload(self):
        self.workbook.save(self.path)
        return self.std.import_code_lists_excel(self.path)

    def code(self, code_list, code):
        return next((entry for entry in self.std.code_lists()[code_list]["codes"] if entry["code"] == code), None)

    def test_the_page_s_own_headings_and_a_meaning_by_its_label_are_read(self):
        renames = {"code_list": "List", "semantic": "Means", "sort_order": "Order", "is_active": "Active"}
        for cell in self.sheet[1]:
            cell.value = renames.get(cell.value, cell.value)
        self.sheet.insert_rows(1)
        self.sheet.cell(1, 1).value = "Code lists, as of today"
        row = self.sheet.max_row + 1
        headings = {cell.value: cell.column for cell in self.sheet[2]}
        for heading, value in {"List": "Failure mechanism", "code": "GIZMO", "name": "Gizmo", "Means": "Unknown"}.items():
            self.sheet.cell(row, headings[heading]).value = value
        outcome = self.upload()
        self.assertEqual([item["code"] for item in outcome["added"]], ["GIZMO"])
        self.assertEqual(self.code("FAILURE_MECHANISM", "GIZMO")["semantic"], "UNKNOWN")

    def test_a_column_twice_refuses_the_workbook(self):
        self.sheet.cell(1, self.sheet.max_column + 1).value = "Code"
        with self.assertRaisesRegex(ValueError, r"^Columns \w+ and \w+ are both code\."):
            self.upload()


if __name__ == "__main__":
    unittest.main()
